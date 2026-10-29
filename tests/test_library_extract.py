import json
import unittest

from bf.documents import load_document
from bf.extract import extract_clauses
from bf.library import Library, LibraryError

from .helpers import LIBRARY, TENDER, TempDirCase, make_pdf


class LibraryTests(TempDirCase):
    def setUp(self):
        super().setUp()
        self.lib = Library.load(LIBRARY)

    def test_people_are_pseudonymised(self):
        rec = self.lib.get("PER-001")
        self.assertEqual(rec.title, "Service Delivery Manager (PER-001)")
        blob = rec.corpus()
        for secret in ("Khalid", "Harbi", "alharbi", "@", "+966", "123 4567"):
            self.assertNotIn(secret, blob)
        self.assertIn("14 years", rec.summary)

    def test_identity_documents_are_never_loaded(self):
        self.assertIsNone(self.lib.get("ID-001"))
        self.assertEqual([a["id"] for a in self.lib.attach_at_submission], ["ID-001"])

    def test_duplicate_ids_rejected(self):
        lib = self.copy_library()
        (lib / "records" / "dup.json").write_text(json.dumps({"id": "CO-001", "type": "company", "title": "x", "summary": "y"}))
        with self.assertRaises(LibraryError):
            Library.load(lib)

    def test_search_prefers_matching_type(self):
        hits = self.lib.search("bid bond bank guarantee", prefer=("financial",))
        self.assertEqual(hits[0][0].id, "FIN-001")

    def test_plurals_match(self):
        cov, shared = self.lib.coverage("bid bonds", "FIN-001")
        self.assertIn("bond", shared)


class ExtractTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.clauses = extract_clauses([load_document(TENDER / "rfp.md")])
        cls.by_text = lambda self, s: next(c for c in self.clauses if s in c.text)

    def test_counts_and_ids(self):
        self.assertEqual(len(self.clauses), 35)
        self.assertEqual([c.id for c in self.clauses][:3], ["C-001", "C-002", "C-003"])

    def test_introduction_without_shall_or_must_is_skipped(self):
        self.assertFalse(any("invites qualified suppliers" in c.text for c in self.clauses))

    def test_flags_and_categories(self):
        iso = self.by_text("ISO 9001")
        self.assertEqual((iso.category, iso.mandatory), ("certification", True))
        desirable = self.by_text("20000-1")
        self.assertFalse(desirable.mandatory)
        self.assertTrue(desirable.scored)
        self.assertEqual(self.by_text("remain valid for 120 days").category, "commercial")
        self.assertEqual(self.by_text("Saudi nationals").category, "nationalisation")
        self.assertEqual(self.by_text("hosted within the Kingdom").category, "data_residency")
        self.assertEqual(self.by_text("Liquidated damages").category, "contract_term")
        self.assertFalse(self.by_text("font size").scored)
        self.assertEqual(self.by_text("ISO 9001").location, "clause 3.1 (Certifications)")

    def test_evaluation_table_rows_become_criteria(self):
        crit = [c for c in self.clauses if c.kind == "criterion"]
        self.assertEqual(len(crit), 6)
        self.assertEqual(sum(c.weight for c in crit), 100)


class FormatTests(TempDirCase):
    def test_docx(self):
        import docx

        d = docx.Document()
        d.add_heading("Certifications", level=2)
        d.add_paragraph("1.1 The supplier must hold a valid ISO 9001 certificate.")
        path = self.tmp / "rfp.docx"
        d.save(str(path))
        clauses = extract_clauses([load_document(path)])
        self.assertEqual(len(clauses), 1)
        self.assertEqual(clauses[0].category, "certification")

    def test_pdf(self):
        path = self.tmp / "rfp.pdf"
        path.write_bytes(make_pdf(["1.1 The supplier must hold a valid ISO 9001 certificate."]))
        doc = load_document(path)
        self.assertIn("ISO 9001", doc.text)
        self.assertEqual(len(extract_clauses([doc])), 1)


if __name__ == "__main__":
    unittest.main()
