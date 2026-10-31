"""Load tender documents as page texts. Supports .txt, .md, .pdf and .docx."""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

# Operator-supplied metadata and the client's pricing sheet are not tender documents.
NON_TENDER_FILES = {"tender.json", "pricing.json"}
SUPPORTED = {".txt", ".md", ".pdf", ".docx"}


class DocumentError(ValueError):
    pass


@dataclass
class Document:
    name: str
    kind: str
    pages: list[str] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n".join(self.pages)


def load_document(path: str | Path) -> Document:
    p = Path(path)
    ext = p.suffix.lower()
    if ext in (".txt", ".md"):
        return Document(p.name, ext[1:], [p.read_text(encoding="utf-8", errors="replace")])
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(p))
        return Document(p.name, "pdf", [(page.extract_text() or "") for page in reader.pages])
    if ext == ".docx":
        import docx  # python-docx

        d = docx.Document(str(p))
        lines = [para.text for para in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                lines.append("| " + " | ".join(cell.text.strip() for cell in row.cells) + " |")
        return Document(p.name, "docx", ["\n".join(lines)])
    raise DocumentError(f"unsupported document type: {p.name}")


def tender_files(tender_dir: str | Path) -> tuple[list[Path], list[str]]:
    """Return (documents to read, names skipped as unsupported)."""
    root = Path(tender_dir)
    if not root.is_dir():
        raise DocumentError(f"tender directory not found: {root}")
    docs, skipped = [], []
    for p in sorted(root.iterdir()):
        if p.name.startswith(".") or not p.is_file() or p.name in NON_TENDER_FILES:
            continue
        if p.suffix.lower() in SUPPORTED:
            docs.append(p)
        else:
            skipped.append(p.name)
    return docs, skipped
